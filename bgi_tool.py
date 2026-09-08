#!/usr/bin/env python3
"""
BGI Script Text Tool
Extracts text from Ethornell/BGI script files into Excel,
and patches translated text back into the scripts.

Supported formats:
  V0           — no header, 16-bit opcodes (older games, SJIS strings)
  V1           — BurikoCompiledScriptVer1.00 header, 32-bit opcodes
  V1-noheader  — 32-bit opcodes, no magic header (newer games, ASCII strings)

Usage:
  Extract:  python bgi_tool.py extract file1 [file2 ...] -o output.xlsx
  Insert:   python bgi_tool.py insert  file1 [file2 ...] -x translations.xlsx -o out_dir/
            python bgi_tool.py insert  file1 [file2 ...] -x translations.xlsx -o out_dir/ --wrapper monospace --line-width 48
"""

import os
import re
import sys
import struct
import argparse
from abc import ABC, abstractmethod
from dataclasses import dataclass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────────────────────
# Word Wrapping
# ─────────────────────────────────────────────────────────────────────────────

LINE_BREAK_CHARS = (' ', '-')
CJK_PUNCTUATION = "，。？！」』】）’”"


class WordWrapper(ABC):
    @abstractmethod
    def get_text_width(self, text: str, offset: int, length: int) -> int:
        pass

    @property
    @abstractmethod
    def line_width(self) -> int:
        pass

    def get_wrap_positions(self, text: str) -> list[int]:
        positions: list[int] = []
        line_start_pos = 0
        n = len(text)

        while line_start_pos < n:
            line_end_pos = line_start_pos
            while line_end_pos < n:
                search_pos = -1
                for idx in range(line_end_pos + 1, n):
                    if text[idx] in LINE_BREAK_CHARS:
                        search_pos = idx
                        break

                if search_pos >= 0:
                    if text[search_pos] != ' ':
                        search_pos += 1
                else:
                    search_pos = n

                if self.get_text_width(text, line_start_pos, search_pos - line_start_pos) > self.line_width:
                    break

                line_end_pos = search_pos

            if line_end_pos == line_start_pos:
                search_pos = line_end_pos
                while line_end_pos < n:
                    search_pos += 1
                    if self.get_text_width(text, line_start_pos, search_pos - line_start_pos) > self.line_width:
                        break
                    line_end_pos = search_pos

            if line_end_pos < n and text[line_end_pos] in CJK_PUNCTUATION:
                line_end_pos += 1

            positions.append(line_end_pos)

            line_start_pos = line_end_pos
            while line_start_pos < n and text[line_start_pos] == ' ':
                line_start_pos += 1

        return positions

    def wrap(self, text: str, line_break: str = "\n") -> str:
        result_lines: list[str] = []
        raw_lines = text.replace("\r\n", "\n").split("\n")

        for line in raw_lines:
            line_start_pos = 0
            wrap_positions = self.get_wrap_positions(line)
            if not wrap_positions:
                result_lines.append("")
                continue

            for line_end_pos in wrap_positions:
                if line_end_pos == line_start_pos:
                    continue
                result_lines.append(line[line_start_pos:line_end_pos])
                line_start_pos = line_end_pos
                while line_start_pos < len(line) and line[line_start_pos] == ' ':
                    line_start_pos += 1

        return line_break.join(result_lines)


class MonospaceWordWrapper(WordWrapper):
    def __init__(self, characters_per_line: int = 50):
        self._characters_per_line = characters_per_line

    @property
    def line_width(self) -> int:
        return self._characters_per_line

    def get_text_width(self, text: str, offset: int, length: int) -> int:
        return max(0, length)


class ProportionalWordWrapper(WordWrapper):
    def __init__(self, font_name: str = "Franklin Gothic Book", font_size: int = 40, bold: bool = False, line_width: int = 1000):
        self._font_name = font_name
        self._font_size = font_size
        self._bold = bold
        self._line_width = line_width
        self._char_widths: dict[str, int] = {}
        self._kern_amounts: dict[tuple[str, str], int] = {}
        self._init_font()

    @property
    def line_width(self) -> int:
        return self._line_width

    def _init_font(self):
        if sys.platform == 'win32':
            try:
                import ctypes
                from ctypes import wintypes

                user32 = ctypes.windll.user32
                gdi32 = ctypes.windll.gdi32

                dc = user32.GetDC(0)
                FW_BOLD = 700
                FW_NORMAL = 400
                ANSI_CHARSET = 0
                weight = FW_BOLD if self._bold else FW_NORMAL

                hfont = gdi32.CreateFontW(
                    self._font_size, 0, 0, 0, weight, 0, 0, 0,
                    ANSI_CHARSET, 0, 0, 0, 0, self._font_name
                )
                gdi32.SelectObject(dc, hfont)

                class ABCFLOAT(ctypes.Structure):
                    _fields_ = [('abcfA', ctypes.c_float), ('abcfB', ctypes.c_float), ('abcfC', ctypes.c_float)]

                class KERNINGPAIR(ctypes.Structure):
                    _fields_ = [('wFirst', wintypes.WORD), ('wSecond', wintypes.WORD), ('iKernAmount', ctypes.c_int)]

                abc_array = (ABCFLOAT * 256)()
                if gdi32.GetCharABCWidthsFloatW(dc, 0, 255, abc_array):
                    for i in range(256):
                        w = int(abc_array[i].abcfA + abc_array[i].abcfB + abc_array[i].abcfC)
                        self._char_widths[chr(i)] = max(0, w)

                num_pairs = gdi32.GetKerningPairsW(dc, 0, None)
                if num_pairs > 0:
                    pairs_array = (KERNINGPAIR * num_pairs)()
                    gdi32.GetKerningPairsW(dc, num_pairs, pairs_array)
                    for pair in pairs_array:
                        self._kern_amounts[(chr(pair.wFirst), chr(pair.wSecond))] = pair.iKernAmount

                user32.ReleaseDC(0, dc)
                gdi32.DeleteObject(hfont)
            except Exception:
                pass

    def _get_char_width(self, c: str) -> int:
        if c in self._char_widths:
            return self._char_widths[c]
        w = self._font_size if ord(c) > 0x2E80 else self._font_size // 2
        self._char_widths[c] = w
        return w

    def _get_kern_amount(self, first: str, second: str) -> int:
        return self._kern_amounts.get((first, second), 0)

    def get_text_width(self, text: str, offset: int, length: int) -> int:
        if length <= 0:
            return 0
        width = 0
        for i in range(offset, offset + length):
            width += self._get_char_width(text[i])
            if i > offset:
                width += self._get_kern_amount(text[i - 1], text[i])
        return width


class NoOpWordWrapper(WordWrapper):
    @property
    def line_width(self) -> int:
        return 999999

    def get_text_width(self, text: str, offset: int, length: int) -> int:
        return 0

    def wrap(self, text: str, line_break: str = "\r\n") -> str:
        return text


# ─────────────────────────────────────────────────────────────────────────────
# Opcode tables
# ─────────────────────────────────────────────────────────────────────────────

V1_MAGIC = b"BurikoCompiledScriptVer1.00\x00"

_V1_WITH_OPS: dict[int, str] = {
    0x0000: "i",    # push constant
    0x0001: "c",    # push code address
    0x0002: "i",    # push var address
    0x0003: "m",    # push string address
    0x0008: "i",
    0x0009: "i",
    0x000A: "i",
    0x0017: "i",
    0x0019: "i",
    0x003F: "i",
    0x007B: "iii",
    0x007E: "i",
    0x007F: "ii",
}

_V1_NO_OPS: set[int] = set()
for _r in [
    range(0x0010, 0x0012), range(0x0015, 0x0020),
    range(0x0020, 0x002C), range(0x0030, 0x0036),
    [0x0038, 0x0039, 0x003A, 0x003E, 0x0040, 0x0048],
    range(0x0080, 0x0084), range(0x0090, 0x009A),
    [0x00A0, 0x00A8, 0x00AA, 0x00AC, 0x00C0, 0x00C1, 0x00C2, 0x00D0],
    range(0x00E0, 0x0500),
]:
    for _x in _r:
        _V1_NO_OPS.add(_x)

V1_ALL_OPS: set[int] = set(_V1_WITH_OPS.keys()) | _V1_NO_OPS
V1_HALT_OPS  = {0x001B, 0x00F4}
V1_FLUSH_OPS = {0x007E, 0x007F, 0x00FE}

V0_OPERAND_TEMPLATES: dict[int, str] = {
    0x0010: "iim", 0x0011: "",   0x0012: "zz",  0x0013: "z",   0x0014: "z",
    0x0015: "",    0x0018: "iiiii", 0x0019: "iiii", 0x001A: "iii",
    0x001B: "ziii",0x001F: "i",  0x0020: "",    0x0021: "",    0x0022: "i",
    0x0024: "iiiii", 0x0025: "ii", 0x0028: "zi", 0x0029: "zzi",
    0x002A: "i",   0x002B: "zi", 0x002C: "ziiiiiiii", 0x002D: "ziiiiiiii",
    0x002E: "iiiii", 0x0030: "zi", 0x0031: "zii", 0x0032: "i",
    0x0033: "i",   0x0034: "ii", 0x0035: "i",   0x0036: "i",   0x0037: "",
    0x0038: "iziiiii", 0x0039: "ii", 0x003A: "iziiiiiiii",
    0x003B: "iiiiii", 0x003C: "iiiiiiiiii", 0x003D: "iiiiiiiiiii",
    0x003F: "i",   0x0040: "iizii", 0x0041: "iizii", 0x0042: "iizi",
    0x0043: "iizi",0x0044: "iizi", 0x0045: "iizi", 0x0046: "izi",
    0x0047: "izi", 0x0048: "ii",  0x0049: "ii",  0x004A: "izi",
    0x004B: "",    0x004C: "zi",  0x004D: "zi",  0x004E: "i",   0x004F: "i",
    0x0050: "zi",  0x0051: "zzi", 0x0052: "i",   0x0053: "zi",  0x0054: "zii",
    0x0060: "iiiii", 0x0061: "ii", 0x0062: "iiiiii",
    0x0065: "i",   0x0066: "ii",  0x0067: "i",   0x0068: "i",   0x0069: "i",
    0x006A: "i",   0x006B: "i",   0x006C: "i",   0x006E: "iii", 0x006F: "i",
    0x0070: "izi", 0x0071: "i",   0x0072: "iii", 0x0073: "iii",
    0x0074: "izi", 0x0075: "i",   0x0076: "iii",
    0x0078: "izi", 0x0079: "i",   0x007A: "iii",
    0x0080: "izii",0x0081: "z",   0x0082: "i",   0x0083: "i",   0x0084: "izi",
    0x0085: "z",   0x0086: "i",   0x0087: "i",   0x0088: "z",
    0x008C: "i",   0x008D: "i",   0x008E: "i",
    0x0090: "i",   0x0091: "i",   0x0092: "i",   0x0093: "i",   0x0094: "i",
    0x0098: "ii",  0x0099: "ii",  0x009A: "ii",  0x009B: "ii",  0x009C: "ii",
    0x009D: "ii",
    0x00A0: "c",   0x00A1: "ic",  0x00A2: "ic",  0x00A3: "iic",
    0x00A4: "iic", 0x00A5: "iic", 0x00A6: "iic", 0x00A7: "iic",
    0x00A8: "iic", 0x00AC: "c",   0x00AD: "",    0x00AE: "i",   0x00AF: "",
    0x00B8: "",    0x00B9: "i",   0x00BA: "i",
    0x00C0: "z",   0x00C1: "z",   0x00C2: "",    0x00C4: "i",
    0x00C8: "z",   0x00C9: "",    0x00CA: "i",
    0x00D0: "",    0x00D4: "i",
    0x00D8: "i",   0x00D9: "i",   0x00DA: "i",   0x00DB: "i",   0x00DC: "i",
    0x00F8: "z",   0x00F9: "zi",  0x00FE: "h",
    0x0110: "zz",  0x0111: "i",
    0x0120: "i",   0x0121: "i",   0x0128: "zii", 0x012A: "ii",
    0x0134: "ii",  0x0135: "i",   0x0136: "i",
    0x0138: "iziiiiziii", 0x013B: "iiiiiiii",
    0x0140: "iiziiii", 0x0141: "iiziiii",
    0x0142: "iiziii",  0x0143: "iiziii",
    0x0144: "iiziii",  0x0145: "iiziii",
    0x0146: "iziii",   0x0147: "iziii",
    0x0148: "ii",  0x0149: "ii",  0x014B: "ziiz",
    0x0150: "zii", 0x0151: "ziii",0x0152: "ii",  0x0153: "iii",
    0x016E: "iiiiii", 0x016F: "iiiiiii", 0x0170: "izzii",
    0x01C0: "zz",  0x01C1: "zz",
    0x0249: "z",   0x024C: "zziii", 0x024D: "z",  0x024E: "zz",  0x024F: "z",
}
V0_SPECIAL_OPS = {0x00A9, 0x00B0, 0x00B4, 0x00FD, 0x0248}
V0_ALL_OPS = set(V0_OPERAND_TEMPLATES.keys()) | V0_SPECIAL_OPS


# ─────────────────────────────────────────────────────────────────────────────
# Format detection
# ─────────────────────────────────────────────────────────────────────────────

def _count_valid_v1_ops(data: bytes, limit: int = 120) -> int:
    pos, count = 0, 0
    while pos + 4 <= len(data) and count < limit:
        op = struct.unpack_from("<I", data, pos)[0]
        if op not in V1_ALL_OPS:
            break
        pos += 4 + 4 * len(_V1_WITH_OPS.get(op, ""))
        count += 1
    return count


def detect_format(data: bytes) -> str:
    if data[: len(V1_MAGIC)] == V1_MAGIC:
        return "v1"
    if _count_valid_v1_ops(data) >= 20:
        return "v1_noheader"
    return "v0"


# ─────────────────────────────────────────────────────────────────────────────
# Data type
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class ScriptString:
    operand_offset: int
    text_offset:    int
    string_type:    str  # 'message' | 'name' | 'internal'


# ─────────────────────────────────────────────────────────────────────────────
# String helpers
# ─────────────────────────────────────────────────────────────────────────────

def _read_sz(data: bytes, offset: int, encoding: str = "shift_jis") -> str:
    end = data.find(b"\x00", offset)
    if end == -1:
        return ""
    raw = data[offset:end]
    for enc in (encoding, "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            pass
    return raw.decode("latin-1", errors="replace")


# ─────────────────────────────────────────────────────────────────────────────
# Punctuation normalization
# ─────────────────────────────────────────────────────────────────────────────

_RE_ELLIPSIS_MARK = re.compile(r"(\.{2,})[ \t]*([?!])\2*")
_RE_REPEATED_MARK = re.compile(r"([!?])\1+")


def _normalize_punctuation(text: str) -> str:
    text = _RE_ELLIPSIS_MARK.sub(lambda m: m.group(2) + m.group(1), text)
    text = _RE_REPEATED_MARK.sub(lambda m: m.group(1) * 3, text)
    return text


# ─────────────────────────────────────────────────────────────────────────────
# Transliteration  (--RU_C / --RU_F)
# ─────────────────────────────────────────────────────────────────────────────

_RU_F_UPPER: dict[str, str] = {
    'A': 'А', 'B': 'Б', 'C': 'В', 'D': 'Г', 'E': 'Д', 'F': 'Е',
    'G': 'Ж', 'H': 'З', 'I': 'И', 'J': 'Й', 'K': 'К', 'L': 'Л',
    'M': 'М', 'N': 'Н', 'O': 'О', 'P': 'П', 'Q': 'Р', 'R': 'С',
    'S': 'Т', 'T': 'У', 'U': 'Ф', 'V': 'Х', 'W': 'Ц', 'X': 'Ч',
    'Y': 'Ш', 'Z': 'Щ', 
}
_RU_F_LOWER: dict[str, str] = {k.lower(): v.lower() for k, v in _RU_F_UPPER.items()}

_RU_F_CYR_TO_BYTES: dict[str, bytes] = {}
for _l, _c in _RU_F_UPPER.items():
    _RU_F_CYR_TO_BYTES[_c] = _l.encode('ascii')
for _l, _c in _RU_F_LOWER.items():
    _RU_F_CYR_TO_BYTES[_c] = _l.encode('ascii')

_RU_F_CYR_TO_BYTES.update({
    'Ъ': b'[',    'Ь': b']',    'ё': b'`',
    'э': b'{',    'ы': b'|',    'я': b'}',
    'Ы': b'\xa1',
    'ь': b'&',
    'ъ': b'+',
    'Ю': b'%',
    'ю': b'$',
    '—': b'#', 
    'Я': b'>',
    'Ё': b'<',
    'Э': b'=',
    'Й': b'J',    'й': b'j',
    '«': b'"',
    '»': b'"',
    '*': b'*',
})


def _apply_ru_f(text: str) -> bytes:
    out = bytearray()
    for ch in text:
        if ch in _RU_F_CYR_TO_BYTES:
            out.extend(_RU_F_CYR_TO_BYTES[ch])
        else:
            try:
                out.extend(ch.encode('cp1252'))
            except (UnicodeEncodeError, ValueError):
                out.extend(b'?')
    return bytes(out)


def _encode_sz_mode(text: str, mode: str) -> bytes:
    if mode == 'direct':
        return text.encode('cp1251', errors='replace') + b'\x00'
    if mode == 'font':
        return _apply_ru_f(text) + b'\x00'
    return text.encode('cp1252', errors='replace') + b'\x00'


# ─────────────────────────────────────────────────────────────────────────────
# V0 extraction
# ─────────────────────────────────────────────────────────────────────────────

def _extract_v0(data: bytes) -> tuple[list[ScriptString], int]:
    strings: list[ScriptString] = []
    pos = 0
    largest_code_addr = 0

    def r_u16() -> int:
        nonlocal pos; v = struct.unpack_from("<H", data, pos)[0]; pos += 2; return v
    def r_i16() -> int:
        nonlocal pos; v = struct.unpack_from("<h", data, pos)[0]; pos += 2; return v
    def r_i32() -> int:
        nonlocal pos; v = struct.unpack_from("<i", data, pos)[0]; pos += 4; return v
    def skip_sz():
        nonlocal pos; end = data.find(b"\x00", pos); pos = (end + 1) if end != -1 else len(data)
    def read_code_addr():
        nonlocal pos, largest_code_addr
        a = struct.unpack_from("<i", data, pos)[0]; largest_code_addr = max(largest_code_addr, a); pos += 4
    def read_str_addr(stype: str):
        nonlocal pos
        off = pos
        a = struct.unpack_from("<i", data, pos)[0]; pos += 4
        is_empty = (0 <= a < len(data) and data[a] == 0)
        strings.append(ScriptString(off, a, "internal" if is_empty else stype))

    def apply_template(tmpl: str):
        for c in tmpl:
            if   c == "h": r_i16()
            elif c == "i": r_i32()
            elif c == "c": read_code_addr()
            elif c == "n": read_str_addr("name")
            elif c == "m": read_str_addr("message")
            elif c == "z": skip_sz()

    while pos < len(data) - 1:
        op_start = pos
        op = r_u16()

        if op == 0x00A9:
            for _ in range(r_i32()): read_code_addr()
        elif op in (0x00B0, 0x00B4):
            for _ in range(r_i32()): skip_sz()
        elif op == 0x00FD:
            for _ in range(r_i32()): skip_sz(); read_code_addr()
        elif op == 0x0248:
            pos = op_start; break
        elif op in V0_OPERAND_TEMPLATES:
            apply_template(V0_OPERAND_TEMPLATES[op])
        else:
            pos = op_start; break

        if op == 0x00C2 and largest_code_addr < pos:
            break

    return strings, pos


# ─────────────────────────────────────────────────────────────────────────────
# V1 extraction
# ─────────────────────────────────────────────────────────────────────────────

def _extract_v1(data: bytes, code_offset: int) -> tuple[list[ScriptString], int]:
    strings: list[ScriptString] = []
    pos = code_offset
    largest_code_addr = 0
    stack: list[tuple[int, int]] = []

    def is_empty(abs_addr: int) -> bool:
        return abs_addr < len(data) and data[abs_addr] == 0

    def flush_as_internal():
        while stack:
            op_off, addr = stack.pop()
            strings.append(ScriptString(op_off, addr, "internal"))

    while pos + 4 <= len(data):
        op = struct.unpack_from("<I", data, pos)[0]
        if op not in V1_ALL_OPS:
            break
        pos += 4

        if op == 0x0003:
            op_off = pos
            addr   = struct.unpack_from("<I", data, pos)[0]; pos += 4
            stack.append((op_off, code_offset + addr))
        elif op == 0x0001:
            a = struct.unpack_from("<I", data, pos)[0]; pos += 4
            largest_code_addr = max(largest_code_addr, code_offset + a)
        elif op == 0x001C:
            if stack:
                op_off, addr = stack.pop()
                strings.append(ScriptString(op_off, addr, "internal"))
                if _read_sz(data, addr) == "_SelectEx":
                    choices, stack[:] = list(reversed(stack)), []
                    for item in reversed(choices):
                        strings.append(ScriptString(item[0], item[1], "message"))
        elif op in (0x0140, 0x0143):
            if stack:
                m_off, m_addr = stack.pop()
                m_type = "message" if not is_empty(m_addr) else "internal"
                if stack:
                    n_off, n_addr = stack.pop()
                    n_type = "name" if not is_empty(n_addr) else "internal"
                    strings.append(ScriptString(n_off, n_addr, n_type))
                strings.append(ScriptString(m_off, m_addr, m_type))
            flush_as_internal()
        elif op == 0x0160:
            choices, stack[:] = list(reversed(stack)), []
            for item in reversed(choices):
                strings.append(ScriptString(item[0], item[1], "message"))
        elif op == 0x0009:
            operand = struct.unpack_from("<I", data, pos)[0]; pos += 4
            if operand == 2 and stack:
                op_off, addr = stack.pop()
                stype = "message" if not is_empty(addr) else "internal"
                strings.append(ScriptString(op_off, addr, stype))
        elif op in V1_FLUSH_OPS:
            flush_as_internal()
            pos += 4 * len(_V1_WITH_OPS.get(op, ""))
        else:
            pos += 4 * len(_V1_WITH_OPS.get(op, ""))

        if op in V1_HALT_OPS and largest_code_addr < pos - code_offset:
            break

    flush_as_internal()
    return strings, pos


def extract_strings(data: bytes) -> tuple[list[ScriptString], int, int, str]:
    fmt = detect_format(data)
    if fmt == "v1":
        code_offset = len(V1_MAGIC) + struct.unpack_from("<I", data, len(V1_MAGIC))[0]
    else:
        code_offset = 0

    if fmt == "v0":
        strings, code_end = _extract_v0(data)
    else:
        strings, code_end = _extract_v1(data, code_offset)

    return strings, code_offset, code_end - code_offset, fmt


# ─────────────────────────────────────────────────────────────────────────────
# Patching
# ─────────────────────────────────────────────────────────────────────────────

def patch_script(data:         bytes,
                 strings:      list[ScriptString],
                 translations: dict[int, tuple[str, str]],
                 code_offset:  int,
                 code_length:  int,
                 mode:         str = "",
                 wrapper:      WordWrapper = None) -> bytes:
    src_enc = "shift_jis" if mode == "" else "cp1252"
    if wrapper is None:
        wrapper = NoOpWordWrapper()

    pool:      bytearray        = bytearray()
    pool_seen: dict[bytes, int] = {}
    new_addrs: dict[int, int]   = {}

    for s in strings:
        if s.operand_offset in translations:
            text, sfx = translations[s.operand_offset]
            
            # ВООБЩЕ НЕ СМОТРИМ НА ТИП СТРОКИ:
            # Если строка содержит пробелы и не помещается в лимит ширины — переносим!
            if len(text) > wrapper.line_width and (' ' in text or '-' in text):
                text = wrapper.wrap(text, line_break="\n")
                
            text = text + sfx
        else:
            text = _read_sz(data, s.text_offset, src_enc)

        encoded = _encode_sz_mode(text, mode)
        if encoded not in pool_seen:
            abs_off = code_offset + code_length + len(pool)
            pool_seen[encoded] = abs_off
            pool.extend(encoded)
        new_addrs[s.operand_offset] = pool_seen[encoded]

    out = bytearray(data[: code_offset + code_length])

    for op_off, new_abs in new_addrs.items():
        struct.pack_into("<i", out, op_off, new_abs - code_offset)

    out.extend(pool)
    return bytes(out)


# ─────────────────────────────────────────────────────────────────────────────
# Excel writing / reading
# ─────────────────────────────────────────────────────────────────────────────

_COL_HEADERS = ["#", "Type", "Original", "TL", "TLE", "Sfx", "Offset (hex)"]
_COL_WIDTHS  = [5,   9,       65,          65,   65,    5,     14]
_HDR_BG, _HDR_FG = "2F5496", "FFFFFF"
_MSG_BG, _NAME_BG = "FFFFFF", "E2EFDA"
_TL_BG  = "FFF2CC"
_TLE_BG = "FCE4D6"


def _border() -> Border:
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def write_xlsx(files: dict[str, tuple[list[ScriptString], bytes, str]],
               out_path: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for filename, (strings, data, encoding) in files.items():
        ws = wb.create_sheet(title=filename[:31])

        for col, (h, w) in enumerate(zip(_COL_HEADERS, _COL_WIDTHS), 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font      = Font(bold=True, color=_HDR_FG, name="Arial", size=10)
            c.fill      = PatternFill("solid", start_color=_HDR_BG)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = _border()
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 18
        ws.freeze_panes = "A2"

        idx = 1
        for s in strings:
            if s.string_type == "internal":
                continue
            text = _read_sz(data, s.text_offset, encoding).replace("\n", "\\n")
            sfx  = ""
            if text.endswith("@"):
                text = text[:-1]
                sfx  = "@"
            row  = idx + 1
            bg   = _NAME_BG if s.string_type == "name" else _MSG_BG

            row_vals = [idx, s.string_type.capitalize(), text, "", "", sfx, hex(s.operand_offset)]
            row_bgs  = [bg,  bg,                         bg,   _TL_BG, _TLE_BG, bg, bg]

            for col, (val, cell_bg) in enumerate(zip(row_vals, row_bgs), 1):
                c = ws.cell(row=row, column=col, value=val)
                c.font      = Font(name="Arial", size=10)
                c.fill      = PatternFill("solid", start_color=cell_bg)
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c.border    = _border()
            idx += 1

    wb.save(out_path)
    print(f"Saved → {out_path}")


def read_xlsx(xlsx_path: str) -> dict[str, dict[int, tuple[str, str]]]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    result: dict[str, dict[int, tuple[str, str]]] = {}
    for ws in wb.worksheets:
        sheet: dict[int, tuple[str, str]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 7:
                continue
            offset_val = row[6]
            tl         = row[3]
            tle        = row[4]
            sfx        = row[5]
            if offset_val is None:
                continue
            try:
                offset = int(str(offset_val), 16)
            except ValueError:
                continue
            sfx_str = str(sfx).strip() if sfx is not None else ""
            translation = None
            if tle is not None and str(tle).strip():
                translation = str(tle)
            elif tl is not None and str(tl).strip():
                translation = str(tl)
            if translation is not None:
                translation = _normalize_punctuation(translation.replace("\\n", "\n"))
                sheet[offset] = (translation, sfx_str)
        result[ws.title] = sheet
    wb.close()
    return result


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────

def collect_files(paths: list[str]) -> list[str]:
    result = []
    for p in paths:
        if os.path.isfile(p):
            result.append(p)
        elif os.path.isdir(p):
            for root, _, files in os.walk(p):
                for f in files:
                    result.append(os.path.join(root, f))
        else:
            print(f"WARNING: not found: {p}")
    return result


def cmd_extract(args) -> None:
    files: dict[str, tuple[list[ScriptString], bytes, str]] = {}
    file_list = collect_files(args.files)

    for filepath in file_list:
        name = os.path.basename(filepath)
        data = open(filepath, "rb").read()
        strings, code_offset, code_length, fmt = extract_strings(data)
        encoding = "shift_jis" if fmt == "v0" else "cp1252"
        visible  = [s for s in strings if s.string_type != "internal"]
        print(f"  {name:30s} format={fmt:<14} "
              f"code=0x{code_offset:05x}+0x{code_length:05x}  "
              f"translatable={len(visible)}")
        files[name] = (strings, data, encoding)

    write_xlsx(files, args.output)


def get_wrapper(args) -> WordWrapper:
    wrapper_type = args.wrapper.lower()
    if wrapper_type == "none":
        return NoOpWordWrapper()
    elif wrapper_type == "monospace":
        return MonospaceWordWrapper(characters_per_line=args.line_width or 50)
    elif wrapper_type == "proportional":
        return ProportionalWordWrapper(
            font_name=args.font_name,
            font_size=args.font_size,
            bold=args.font_bold,
            line_width=args.line_width or 1000
        )
    else:
        raise ValueError(f"Unknown wrapper type: {wrapper_type}")


def cmd_insert(args) -> None:
    translations_by_file = read_xlsx(args.xlsx)
    os.makedirs(args.output, exist_ok=True)
    file_list = collect_files(args.files)

    if args.RU_C:
        mode = 'direct'
        mode_label = 'RU_C (cp1251 кириллица)'
    elif args.RU_F:
        mode = 'font'
        mode_label = 'RU_F (специализированный шрифт)'
    else:
        mode = ''
        mode_label = 'ASCII/cp1252'

    wrapper = get_wrapper(args)
    print(f"Режим: {mode_label}")
    if args.wrapper == "none":
        print("Перенос строк: отключен")
    elif args.wrapper == "monospace":
        print(f"Перенос строк: моноширинный ({wrapper.line_width} символов)")
    else:
        print(f"Перенос строк: пропорциональный ({args.font_name} {args.font_size}px, макс {wrapper.line_width}px)")

    for filepath in file_list:
        name = os.path.basename(filepath)
        data = open(filepath, "rb").read()
        strings, code_offset, code_length, fmt = extract_strings(data)

        sheet_key = name[:31]
        if sheet_key not in translations_by_file:
            print(f"  WARNING: no sheet '{sheet_key}' found in Excel — skipping {name}")
            continue

        trans = translations_by_file[sheet_key]
        patched  = patch_script(data, strings, trans, code_offset, code_length, mode, wrapper)
        out_path = os.path.join(args.output, name)
        with open(out_path, "wb") as f:
            f.write(patched)
        print(f"  Patched → {out_path}  ({len(trans)} strings replaced)")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract/insert text in Ethornell BGI script files.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры:
  python bgi_tool.py extract scripts/ -o strings.xlsx
  python bgi_tool.py insert  scripts/ -x strings.xlsx -o patched/ --wrapper monospace --line-width 48
  python bgi_tool.py insert  scripts/ -x strings.xlsx -o patched/ --RU_C --wrapper monospace --line-width 48
  python bgi_tool.py insert  scripts/ -x strings.xlsx -o patched/ --RU_F --wrapper monospace --line-width 48
""",
    )
    sub = parser.add_subparsers(dest="command")

    p_ext = sub.add_parser("extract", help="Извлечь текст в Excel")
    p_ext.add_argument("files", nargs="+", help="Файлы или папки со скриптами")
    p_ext.add_argument("-o", "--output", default="strings.xlsx", help="Выходной .xlsx")

    p_ins = sub.add_parser("insert", help="Вставить переводы обратно в скрипты")
    p_ins.add_argument("files", nargs="+", help="Оригинальные файлы или папки")
    p_ins.add_argument("-x", "--xlsx", required=True, help="Excel с переводами")
    p_ins.add_argument("-o", "--output", default="output", help="Выходная папка")

    mode_group = p_ins.add_mutually_exclusive_group()
    mode_group.add_argument("--RU_C", action="store_true", help="Прямой вывод кириллицы (cp1251)")
    mode_group.add_argument("--RU_F", action="store_true", help="Специализированный шрифт")

    p_ins.add_argument("--wrapper", choices=["monospace", "proportional", "none"], default="monospace",
                       help="Алгоритм переноса (по умолчанию: monospace)")
    p_ins.add_argument("--line-width", type=int, default=None,
                       help="Ширина строки: символов для monospace (по умолчанию 50), пикселей для proportional (по умолчанию 1000)")
    p_ins.add_argument("--font-name", default="Franklin Gothic Book", help="Имя шрифта для proportional")
    p_ins.add_argument("--font-size", type=int, default=40, help="Размер шрифта для proportional")
    p_ins.add_argument("--font-bold", action="store_true", help="Жирный шрифт для proportional")

    args = parser.parse_args()
    if args.command == "extract":
        cmd_extract(args)
    elif args.command == "insert":
        cmd_insert(args)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()